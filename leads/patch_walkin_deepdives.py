# -*- coding: utf-8 -*-
"""Apply Tue-15 deep-dive findings to leads_50.xlsx."""
import openpyxl

wb = openpyxl.load_workbook('leads/leads_50.xlsx')
ws = wb['Leads 50']
hdr = [c.value for c in ws[1]]
C = {h: i + 1 for i, h in enumerate(hdr)}

def setrow(lid, **kw):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == lid:
            for k, v in kw.items():
                key = k.replace('_', ' ') if k not in C else k
                ws.cell(r, C[key]).value = v
            return r
    raise KeyError(lid)

def append_note(lid, note):
    r = setrow(lid)
    c = ws.cell(r, C['Notes'])
    c.value = (c.value or '') + ' ' + note

append_note(18, '| DEEP DIVE 14 Sep: 652 075 229 = official line self-listed on GPENreformation (confirmed on WA King 14 Sep = school line, not principal personal). 675 533 321 = Kinang Edwin Ngenge, staff/language teacher & network partnership contact (NOT decision-maker; internal champion). Sister school PCSS Bonamoussadi Douala already runs pcssbonamoussadi.online with WhatsApp = network precedent (do not name). Boarding 700+ = headline hook.')
append_note(20, '| DEEP DIVE 14 Sep: located near Baptist church (Ndongo) and CUIB, Molyko; technical A-Level 91.7% pass 2016 incl Accounting; GCE centres 11412 + 22089, hosts external candidates. Zero online contacts = walk-in only; confirm faith affiliation + enrollment at gate.')
append_note(28, '| DEEP DIVE 14 Sep: LOCATION Check Point (Wokoko), opposite ENAMEN Pharmacy. Principal Mr Yerima Samson Tata (Cameroon Tribune Jun/Sep 2026). LIVE PAIN: principal publicly reported slower 2026/27 enrollment (parents arriving Oct) while spending on repainting/benches. Facilities: science lab, computer lab + internet, band, generator, reservoir; UNIQUE evening school 15:30. Mission text commits to technology + bilingualism. ~1,350 GCE candidates hosted 2026.')
append_note(29, '| DEEP DIVE 14 Sep: located former HIMS Campus, Garden Park Molyko, opposite Orange Cameroon; motto Knowledge-Fidelity-Integrity; founder/proprietor Dr Bernard Fomba (ask by name, decides alone); 134 candidates at one 2025 session.')
append_note(30, '| DEEP DIVE 14 Sep: school gives its name to Nabesk junction, Bonduma (press landmark Mar 2025); 83.5% O-level 2020; external candidate host (214/187/401 in 2025 movement lists). Ask who founded it (name may derive from proprietor).')
setrow(31,
       Website='bairdmemorial.com (self-built, spelling errors; unreachable from outside 14 Sep - verify on phone)',
       **{'Website status': 'Site EXISTS (weak/D-) - CONDITIONAL pitch'},
       Phone='+237 677 87 53 95 / 677 78 04 05 / 99 83 74 79',
       WhatsApp='check 677 87 53 95 + 677 78 04 05',
       Decision_maker='Principal/proprietor (capture); P.O. Box 403 Buea',
       Lead_score=6, Priority='C conditional',
       Notes=("Walk-In-Batch-2026-09-15 stop 6 light. Boarding 'inclusive' high school behind public tap Bonduma; bairdmemorial@yahoo.com; GCE 11489 (small); active CSDC inter-school events 2025. "
              "| DEEP DIVE 14 Sep: ALREADY HAS self-built website bairdmemorial.com (visible typos 'bording/Accademic', did not open externally). Phone-check tonight: DEAD = full D pitch ('I tried opening it before coming'); WORKS = card + soft refresh/upsell line, never criticise at gate."))
append_note(32, '| DEEP DIVE 14 Sep: 679 65 07 07 confirmed on WA (King) — attribution = number published on the college FB page intro (official school line, role unknown; likely admin/proprietor; verify profile photo/name tonight). Baptist confessional -> principal then Baptist education office; ask who else decides. bgccbuea@yahoo.com.')

new = [
    {"ID": 34, "School": "Bishop Jules Peters Memorial College", "City": "Buea (Bokwaongo)", "Language": "EN",
     "Facebook": "N/V", "Website": "None found", "Website status": "D/E - no site (diocese education pages thin)",
     "Admissions activity": "Yes - Inaugural Mass 2025/26 held 17 Feb 2026",
     "Phone": "via diocese Bishop's House Small Soppo +237 683 348 856", "WhatsApp": "verify diocesan number",
     "Decision maker": "Proprietor = Bishop Michael Bibi; principal is a woman (name at gate); Education Secretariat ratifies",
     "Contact channel": "WALK-IN Tue 15 stop 8 (replaces disqualified Inter Comp)",
     "Facilities": "GCE centre 11535 (small, ~35 candidates 2024); young/small diocesan college",
     "Lead score": 9, "Priority": "B",
     "Notes": "Walk-In-Batch-2026-09-15. Catholic diocesan pipeline (same motion as Sasse): principal first, offer forwardable preview, ask when Education Secretariat next meets; multi-college upside. Bokwaongo uphill past Great Soppo. Source: bueadiocese.org 18 Feb 2026."},
    {"ID": 35, "School": "Inter Comprehensive High School (ICHS) Great Soppo", "City": "Buea (Great Soppo)",
     "Language": "EN", "Website": "None found", "Website status": "D/E - no site",
     "Phone": "not listed", "WhatsApp": "n/a",
     "Decision maker": "DISPUTED - founder Enni Philomena vs administrators of late George Fongoh Mayah",
     "Contact channel": "DO NOT APPROACH - legal/management dispute",
     "Facilities": "GCE centre 11216 (~72 candidates); opened 2024/25 amid public ownership dispute",
     "Lead score": 0, "Priority": "X disqualified",
     "Notes": "Disqualified 14 Sep (Walk-In-Deep-Dives): public ownership lawsuit reported mimimefoinfos.com 25 Sep 2024 (school begun as Stenographics venture, became Inter-Comprehensive College 1991). No one can sign. Re-check after ~6 months of stable single management."},
]
for nl in new:
    ws.append([nl.get(h) for h in hdr])

# Pipeline tab rows
ps = wb['Pipeline - 5 Stages']
h2 = [c.value for c in ps[1]]
def pcol(slug):
    for i, h in enumerate(h2, start=1):
        if h and slug.lower() in str(h).lower():
            return i
    return None
ci_name, ci_stage, ci_next = pcol('lead') or 2, pcol('stage') or 6, pcol('next') or 7
for nl in new:
    ps.append([None] * (ci_name - 1) + [nl['School']] + [None] * (ci_stage - ci_name - 1) +
              ['Walk-in Tue 15' if nl['ID'] == 34 else 'DISQUALIFIED',
               ('Bokwaongo stop 8; principal-first diocesan motion; forwardable preview; concept within 24h on warm yes.'
                if nl['ID'] == 34 else
                'Ownership dispute; revisit ~Mar 2027 if single management confirmed.')])

wb.save('leads/leads_50.xlsx')
print('patched; leads now', ws.max_row - 1)
