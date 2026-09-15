# -*- coding: utf-8 -*-
"""15 Sep correction: Solidarity leadership/numbers per RESEARCH-STANDARD; add One Stop lab lead."""
import openpyxl, pathlib

P = pathlib.Path(__file__).resolve().parent / "leads_50.xlsx"
wb = openpyxl.load_workbook(P)
ws = wb["Leads 50"]
hdr = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(hdr)}

for row in ws.iter_rows(min_row=2):
    if row[col["ID"]].value == 36:
        vals = {
            "Decision maker": "CURRENT DIRECTOR UNKNOWN publicly (gatekeeper routing). Founder Dr Peter Nde Fon (CEO) DIED 21 Oct 2020 (Cameroon-Info/MMI obits); run by Solidarity Health Foundation CIG, family incl. widow Beatrice Nde Fon plausibly involved (unconfirmed). Current MO: Dr Asoba Anodem (LinkedIn, 2011-present). Dr Njang Mbeng Emmanuel listed on stale medicoor roster but NOT confirmed current (his LinkedIn: DMO Muyuka + co-owner One Stop lab; do not assume)",
            "Phone": "CLINIC LINE 677 61 57 57 (ORG-LIKELY: GBP+elidge+medicoor; King WA-profile check tonight); landline 2 33 32 31 31; admin candidate 677 61 12 07 (UN NGO registry, unverified)",
            "WhatsApp": "VERIFY 677 61 57 57 as clinic line tonight. DO NOT TEXT 691 63 29 41 about Solidarity: King confirms it is Dr Njang's PERSONAL WhatsApp; his current role at the clinic is unconfirmed and he co-owns a competing lab",
            "Contact channel": "Tue 15 Sep 08:35 WhatsApp msg 1 (Clinic-Batch section 1, NO personal names) to clinic line 677 61 57 57 after King's WA-profile check; asks for current medical director OR clinic manager. Yes = nameless /sample-clinic.html + invite choice; named within 24h",
            "Notes": "Clinic lead #1, research-standard corrected 15 Sep (see sales/RESEARCH-STANDARD.md error #1). Not-for-profit CIG founded Aug 1998 by late Dr Peter Nde Fon (UB public health chair, obituary verified). GBP 4.0 stars/54 reviews/21 photos/24h, Untarred Malingo St plus code 575J+7M, NO website ('Add website'), medicoor unclaimed+stale (still lists deceased founder), X @SolidarityHeal1 since May 2021. Services: gen med, vaccination, surgery, lab panels, maternity, ultrasound/X-ray; reviews cite cardiologist+pharmacy; Zenithe accredited. Among Buea district top-3 most-solicited facilities (PAMJ 2023, BMRI 2025). Pitch: attach bilingual site+WA booking+same-day lab results to GBP; foundation/grant visibility; claim medicoor; 24h emergency bar. FUs M+2/M+4/M+7.",
        }
        for k, v in vals.items():
            row[col[k]].value = v
        break
else:
    raise SystemExit("ID 36 missing")

existing = {row[col["ID"]].value for row in ws.iter_rows(min_row=2)}
if 38 not in existing:
    ws.append([
        38, "One Stop Medical Laboratory & Diagnostics", "Buea (location TBD; co-owner based Buea)", "Bilingual EN/FR",
        "not researched yet", "not found yet",
        "Unknown - research pending (founded May 2023 per owner LinkedIn)",
        "n/a", "n/a", "medical laboratory + diagnostics (co-owned with a second owner; scope TBD)",
        "+237 691 63 29 41 (PERSONAL line of co-owner Dr Njang Mbeng Emmanuel per King; WA confirmed personal, do not use as org front door)",
        "691 63 29 41 = Dr Njang personal (King); ORG number unknown",
        "Dr Njang Mbeng Emmanuel MD MBA MSc (co-owner; also DMO Muyuka Health District Oct 2020-present; ex Maealth mobile-health Medical Director; FAMSA ex-president; UBMSAA president; HERO Cameroon CFO)",
        "HOLD - do not contact until (1) Solidarity relationship clarified and (2) dossier researched per RESEARCH-STANDARD. High digital-health fit (self-described 'digital driven healthcare solutions', built Help Yourself app). Separate dossier required; never mix with Solidarity pitch",
        "laboratory and diagnostics", "unknown",
        12, "B+ (potential)", "none yet", "not contacted", "", "", "", "", None, "research then schedule",
        "Created 15 Sep during Solidarity deep recheck. Entrepreneurial public-health doctor with proven appetite for digital health = strong future website/app client. Conflict-aware: he appears on Solidarity's stale medicoor roster; resolve that first. Follow-ups M+2/M+4/M+7 once opened.",
    ])

wb.save(P)
print("saved, rows:", ws.max_row)
