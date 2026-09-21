# -*- coding: utf-8 -*-
"""St Theresa STIBCCOL — first-ever reply logged (Tue 15 Sep 20:44):
soft objection 'website already in the making, ready by October, too late'.
Soft-ack branch: congratulate, no pitch, plant two must-haves, ask
permission to check back at launch; FU 14 Oct 2026. TikTok fallback CANCELLED.
NOTE: this chat has 24h disappearing messages — log verbatim here."""
import openpyxl, pathlib

P = pathlib.Path("leads/leads_50.xlsx")
wb = openpyxl.load_workbook(P)
ws = wb["Leads 50"]
hdr = [c.value for c in ws[1]]
idx = {h: i + 1 for i, h in enumerate(hdr)}
ROW = 28

ws.cell(row=ROW, column=idx["Contacted"],
        value="YES 15 Sep 15:45 WA by King (text-only msg 1, 679 15 10 75)")
ws.cell(row=ROW, column=idx["Reply"],
        value="YES Tue 15 Sep 20:44 (verbatim): 'Gd evening we already have one in the making which will be ready by October. If you had gotten to me earlier then we could have used your system but it's already late.'")
ws.cell(row=ROW, column=idx["Conversation"],
        value="SOFT OBJECTION = competitor/in-house build in progress, ETA October. Soft-ack reply sent by King same evening: congrats, fees+WhatsApp-admissions checklist, free second-opinion offer if it slips, permission to check back at launch. NO mockup/link pushed (would damage). 24h disappearing messages on this chat.")
ws.cell(row=ROW, column=idx["Follow-up date"],
        value="Wed 14 Oct 2026 (site promised for October — check-in: is it live? free honest review incl. WA admissions/bilingual/fees/3G); FU2 Wed 4 Nov if silent; then stop. No TikTok DM fallback (engaged on WA).")
ws.cell(row=ROW, column=idx["Notes"],
        value=(ws.cell(row=ROW, column=idx["Notes"]).value or "") +
        " | 15 Sep 20:44 FIRST REPLY OF THE CAMPAIGN: site 'in the making, ready October, already late'. Soft-ack branch only. Watch the October launch (check Google/FB): template/student builds commonly miss WhatsApp admissions, bilingual toggle, fees pages and 3G speed = re-open wedge. Verified GCE Board 2026 marking centre (camgceb.org) — usable compliment at FU.")

# Daily Tracker — count the reply on 15 Sep
tr = wb["Daily Tracker"]
th = [c.value for c in tr[1]]
ti = {h: i + 1 for i, h in enumerate(th)}
done = False
for row in tr.iter_rows(min_row=2):
    v = row[ti["Date"] - 1].value
    if v and "2026-09-15" in str(v):
        c = row[ti["Replies"] - 1]
        c.value = (c.value or 0) + 1 if isinstance(c.value, (int, float)) else 1
        c = row[ti["Positive replies"] - 1]
        # courteous reply, kept door open = warm-ish; counted as reply not positive-buy
        n = row[ti["Notes"] - 1]
        n.value = (n.value or "") + " | St Theresa replied 20:44 (site in making, Oct) — soft-ack sent; first inbound reply of campaign."
        done = True
        break
if not done:
    print("WARN: no 2026-09-15 tracker row found")

wb.save(P)
print("patched", P)
