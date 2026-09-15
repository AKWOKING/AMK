# -*- coding: utf-8 -*-
"""15 Sep deep recheck of Solidarity: corrected doctor name + Google profile facts."""
import sys
raise SystemExit("SUPERSEDED 15 Sep: Solidarity facts corrected by patch_solidarity_leadership.py (founder deceased; Njang unconfirmed). Do not rerun this script. Kept for history.")
import openpyxl, pathlib

P = pathlib.Path(__file__).resolve().parent / "leads_50.xlsx"
wb = openpyxl.load_workbook(P)
ws = wb["Leads 50"]
hdr = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(hdr)}

for row in ws.iter_rows(min_row=2):
    if row[col["ID"]].value == 36:
        vals = {
            "School": "Solidarity Health Foundation (Solidarity Clinic & Laboratory)",
            "City": "Buea (Untarred Malingo St / Molyko Checkpoint D61, P.O. Box 467; plus code 575J+7M)",
            "Website": "none attached - Google profile shows 'Add website'; medicoor profile unclaimed",
            "Website status": "NO website, but RICH Google Business Profile: 4.0 stars / 54 reviews / 21 photos, open 24h, 677 61 57 57 - goodwill dead-ends at a phone number",
            "Admissions activity": "24h clinic: general medicine, vaccination, surgery, laboratory (hematology/parasitology/serology/microbiology/biochem/hormonology), maternity, ultrasound/X-ray; reviews cite cardiologist, pharmacy, free specialist consults",
            "Phone": "+237 691 63 29 41; +237 677 61 57 57 (King: number correct); landline 2 33 32 31 31",
            "WhatsApp": "KING TO VERIFY 691 63 29 41 and 677 61 57 57 on WA (677 cross-confirmed by Google/medicoor/elidge/WTWC/Zenithe)",
            "Decision maker": "Dr Nde Fon Peter (GP per medicoor staff list; WTWC 'NDIFOR' = mis-transcription, The Sun renders 'Ndifon'). Second GP: Dr Njang Mbeng Emmanuel. NOT Dr Peter Louis Ndifor of Bota Polyclinic Limbe (different man)",
            "Contact channel": "Tue 15 Sep 08:35 WhatsApp msg 1 (Clinic-Batch section 1, KING PRIORITY): names Google 4.0/54 reviews with no website/booking; asks which doctor receives it. Yes = nameless /sample-clinic.html + invite choice; named within 24h",
            "Lead score": 15,
            "Notes": "Clinic lead #1, upgraded after deep recheck 15 Sep. Social-enterprise foundation; among Buea district's 3 most-solicited facilities (Pan African Med J 2023 + BioMed Research Intl 2025). Zenithe insurance accredited (independently confirmed); medicoor shows big insurer panel (verify before citing). Pitch: attach bilingual site + WA booking + lab results-on-WA to the existing Google profile; claim medicoor; 24h emergency bar. Follow-ups M+2/M+4/M+7. Ignore solidarity-clinic.blogspot.gr (Athens Greece 2013, unrelated).",
        }
        for k, v in vals.items():
            row[col[k]].value = v
        break
else:
    raise SystemExit("ID 36 missing")

wb.save(P)
print("saved")
