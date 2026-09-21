# -*- coding: utf-8 -*-
"""15 Sep afternoon: St Theresa SENT (1 tick); Baird GO after personal-line +
dead-domain recheck; mockup-image standing rule; MITOC FU schedule in channel."""
import openpyxl

wb = openpyxl.load_workbook('leads/leads_50.xlsx')
ws = wb['Leads 50']
hdr = [c.value for c in ws[1]]
C = {h: i + 1 for i, h in enumerate(hdr)}

def upd(lid, **kw):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == lid:
            for k, v in kw.items():
                key = k if k in C else k.replace('_', ' ')
                ws.cell(r, C[key]).value = v
            return r
    raise KeyError(lid)

def note(lid, text):
    r = upd(lid)
    c = ws.cell(r, C['Notes'])
    c.value = (c.value or '') + ' ' + text

# 27 St Theresa — King sent msg 1 himself before the mockup rule existed
upd(27, **{'Contact channel':
    'SENT by King ~15:45 Tue 15 Sep (one grey tick; text-only). No chase; TikTok @stibccol DM = Wed second touch only if WA stays silent. On warm reply: attach demos/shots/mockup-secondary-wa.jpg + live /sample-secondary.html link, then Thu/Fri walkthrough choice'})
note(27, '| 15 Sep PM: msg 1 out (1 tick, not yet delivered to handset). Personal WA line (lay private = likely proprietor pocket = buyer). Mockup image went to later schools first; include image with the link on reply.')

# 31 Baird — GO granted; pack prepped; King has not confirmed the tap yet
upd(31,
    **{'Website status': "bairdmemorial.com DNS-DEAD re-verified 15 Sep 13:35 (apex+www NXDOMAIN; Google still indexes the dead Home + /about-us -> say 'leads parents nowhere', never 'finds nothing')"},
    **{'Contact channel':
    "GO (King 'proceed' 15 Sep PM): attach demos/shots/mockup-secondary-wa.jpg FIRST, then recipient-neutral msg (replan §4) to WA 677 87 53 95; pre-tap open bairdmemorial.com once, swap para 3 if it loads; Wed sealed card in Limbe sweep if silent; invited visit only"})
note(31, "| 15 Sep PM GO: 95 line is a PERSONAL WA account (man's profile photo, no business profile); their OWN old dead site names PROPRIETRESS Madam Mary Forju (also Mr Ndichafah Fredrick, extra nums 652 24 80 21/677 16 98 66/673 40 40 69, gmail/outlook/yahoo boxes - record only, never open). Holder of 95 unknown (principal/bursar/family possible) -> msg asks 'proprietor or the principal?', never names her, never references the photo. Never mock old copy; ask who built it/last touched on warm reply.")

# 37 MITOC — channel line carries the exact FU cadence now
upd(37, **{'Contact channel':
    'SENT msg 1 Tue (two ticks, awaiting reply). No chase same day; FU1 Thu 17 / FU2 Sat 19 / FU3 Tue 22 Sep (OWNER-BUYER only, then stop). Warm yes -> https://mitoc-concept.vercel.app + 10-min shop visit vs all-on-WhatsApp; amk-cm /mitoc.html stays 404 until King redeploys amk-site.zip'})

wb.save('leads/leads_50.xlsx')
print('xlsx patched: 27, 31, 37')
