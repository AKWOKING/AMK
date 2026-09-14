# -*- coding: utf-8 -*-
"""15 Sep update: King's WA verifications + clinic batch (Solidarity, MITOC).
Idempotent: run repeatedly, rows matched by ID."""
import openpyxl, pathlib

P = pathlib.Path(__file__).resolve().parent / "leads_50.xlsx"
wb = openpyxl.load_workbook(P)
ws = wb["Leads 50"]
hdr = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(hdr)}

def set_row(rid, **vals):
    for row in ws.iter_rows(min_row=2):
        if row[col["ID"]].value == rid:
            for k, v in vals.items():
                row[col[k]].value = v
            return True
    raise SystemExit(f"ID {rid} not found")

# --- King's 14 Sep confirmations on existing rows ---
set_row(27,
        WhatsApp="CONFIRMED on WA King 14 Sep: 679 15 10 75",
        **{"Contact channel": "Tue 15 Sep 15:45 WhatsApp msg 1 (replan §4); TikTok @stibccol DM Wed only if WA silent (no same-morning double touch)"})

set_row(28,
        **{"Contact channel": "WED 16 sealed-envelope card drop to Principal Tata = ONLY counted first touch. Tue email REMOVED per King (emails rarely opened); smbicol@ bonus only, never planned on"})

set_row(31,
        Phone="+237 677 87 53 95",
        **{"Website status": "bairdmemorial.com DEAD - confirmed on King's phone 14 Sep (full pitch, not refresh)",
           "WhatsApp": "CONFIRMED on WA King 14 Sep: 677 87 53 95. 677 78 04 05 NOT on WhatsApp - DISCARDED",
           "Contact channel": "Tue 15 Sep 16:00 WhatsApp msg 1 FULL pitch (replan §4); if silent, sealed card in Wed Limbe sweep; invited Bonduma visit only after",
           "Lead score": 11,
           "Priority": "B+"})

set_row(32,
        WhatsApp="CONFIRMED on WA King 14 Sep: 679 65 07 07",
        **{"Contact channel": "Tue 15 Sep 09:00 WhatsApp msg 1 (replan §4) then nameless sample link + invite-choice walkthrough Thu/Fri; Baptist protocol: principal then education office"})

# --- New clinic leads ---
existing = {row[col["ID"]].value for row in ws.iter_rows(min_row=2)}
new_rows = [
 [36, "Solidarity Clinic & Laboratory", "Buea (Molyko, P.O. Box 467)", "Bilingual EN/FR",
  "none (unclaimed medicoor.com profile id 380, no photo/ratings)", "none",
  "No website - only a bare third-party booking profile; cross-listed in WTWC PDF and Zenithe Insurance accredited providers as 'Solidarity Health Foundation, Molyko'",
  "n/a", "n/a", "clinic + lab: hematology/parasitology/serology/biochemistry/microbiology, ultrasound, X-ray, consultations",
  "+237 691 63 29 41; +237 677 61 57 57", "KING TO VERIFY 691 63 29 41 and 677 61 57 57 on WA (677 cross-confirmed by 3 sources)",
  "Dr Ndifor (per WTWC provider PDF; role attribution at gate)",
  "Tue 15 Sep 08:35 WhatsApp msg 1 (Clinic-Batch §1, KING PRIORITY) -> yes = nameless /sample-clinic.html + invite choice (Thu/Fri walkthrough vs WhatsApp); named within 24h",
  "laboratory + ultrasound/X-ray + insurance accreditation", "no",
  14, "A", "nameless sample-clinic.html concept built 15 Sep (Forest palette)", "scheduled Tue", "", "", "", "", None, "Tue 15 Sep",
  "Clinic lead #1. Pitch hooks: results on WhatsApp same day, transparent FCFA lab panel prices, insurance-accredited invisible on Google, EN|FR, emergency line. Possibly foundation/NGO per accreditation doc - speak to the doctor. Sources: medicoor profile 380, WTWC provider PDF, Zenithe list. Follow-ups M+2/M+4/M+7."],
 [37, "Midas Touch Optic Center (MITOC)", "Buea (Molyko, opp former police station, Malingo)", "Bilingual EN/FR",
  "Facebook page id 100064126300520 (~365 likes, Medical & health); 'We refract, prescribe n mount lenses'", "none",
  "Facebook-only; no website, no Google listing; intro carries no hours/prices",
  "n/a", "365 likes", "refraction/eye tests, prescribing + mounting lenses (full optician, not frame shop)",
  "+237 678 90 89 62", "KING TO VERIFY 678 90 89 62 on WA (number from official FB intro)",
  "Ateasom Collins (name associated with FB page; proprietor/optician to confirm)",
  "Tue 15 Sep 08:35 WhatsApp msg 1 (Clinic-Batch §2, KING PRIORITY) -> yes = nameless /sample-clinic.html (optic re-skin promised in 24h) + invite choice",
  "eye testing + lens mounting; frame catalog content he supplies", "no",
  12, "B+", "nameless sample-clinic.html concept (re-skin to optic) built 15 Sep", "scheduled Tue", "", "", "", "", None, "Tue 15 Sep",
  "Clinic lead #2. Pitch hooks: invisible on Google (FB-only), eye-test booking on WA, frames + FCFA lens prices bilingual, student catchment Molyko. Smaller scope than Solidarity - quote by scope, never discount. Follow-ups M+2/M+4/M+7."],
]
for nr in new_rows:
    if nr[0] not in existing:
        ws.append(nr)

wb.save(P)
print("saved", P, "rows:", ws.max_row)
