# -*- coding: utf-8 -*-
"""Invitation-first replan: channel updates for walk-in batch."""
import openpyxl

wb = openpyxl.load_workbook('leads/leads_50.xlsx')
ws = wb['Leads 50']
hdr = [c.value for c in ws[1]]
C = {h: i + 1 for i, h in enumerate(hdr)}

def upd(lid, **kw):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == lid:
            for k, v in kw.items():
                key = k if k in C else k.replace('_', ' ')
                ws.cell(r, C[key]).value = v
            return
    raise KeyError(lid)

def note(lid, text):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == lid:
            c = ws.cell(r, C['Notes'])
            c.value = (c.value or '') + ' ' + text
            return
    raise KeyError(lid)

# 27 St Theresa
upd(27,
    Phone='+237 679 15 10 75 (official FB page intro; King to WA-verify); secondary 677 36 19 11 low confidence',
    WhatsApp='CHECK 679 15 10 75 tonight',
    **{'Contact channel': 'Tue 15: TikTok DM @stibccol (active to 14 Sep 2026) + WhatsApp 679 15 10 75 if confirmed'})
note(27, '| INVITATION-FIRST 14 Sep: TikTok @stibccol ACTIVE (Bilingualism Day/Miss St Theresa posts 14 Sep 2026) = DM channel like Sasse; official FB page lists 679 15 10 75 (verify WA profile tonight). YouTube 2022 "for more info" 677 36 19 11 likely media person, do not pitch.')

# 28 Summerset
upd(28,
    Facebook='Page facebook.com/summerset.bilingual.college INACTIVE since Dec 2015 (stray email dadariodelanuez@yahoo.com - do not use)',
    Phone='not published anywhere',
    WhatsApp='none found',
    **{'Contact channel': 'Tue parallel email smbicol@yahoo.com (weak); Wed async card-drop addressed to Principal Tata; invited visit only after'})
note(28, '| INVITATION-FIRST 14 Sep: NO reachable digital channel - no phone online, FB dead since 2015. Plan: email Tue + sealed-envelope card drop Wed (hand-addressed to Principal Yerima Samson Tata, opposite ENAMEN Pharmacy); QR inbound = they initiate; visit by invitation only.')

# 31 Baird
note(31, '| INVITATION-FIRST 14 Sep: WA-check 677 87 53 95 + 677 78 04 05 profile attribution tonight; message only if site dead AND number on WA (text in replan section 4); otherwise Wed card drop same road as NABESK.')

# 20 Salvation, 29 SBHS, 30 NABESK, 33 St Sylvester -> card-drop sweep
for lid, who in [(20, 'Salvation Molyko near CUIB'), (29, 'SBHS Garden Park opposite Orange, attn Dr Fomba'),
                 (30, 'NABESK Nabesk junction'), (33, 'St Sylvester Muea (fold in only if Muea trip)')]:
    upd(lid, **{'Contact channel': f'Wed 16 async card-drop sweep ({who}); no meeting, sealed envelope to principal; visit only by invitation'})

# 32 Baptist + 18 PCSS channels set
upd(32, **{'Contact channel': 'Tue 08:30 WhatsApp 679 65 07 07 (CONFIRMED on WA King 14 Sep; school line) then preview link; invite-choice walkthrough Thu/Fri; gate visit only if invited'})
upd(18, **{'Contact channel': 'Tue 08:30 WhatsApp 652 075 229 (CONFIRMED on WA; official line); Kinang 675 533 321 door-opener at noon if silent; visit by invitation'})

# 34 Bishop Jules
upd(34, **{'Contact channel': 'Hold for diocesan motion behind Sasse; fold card/visit into Great Soppo/Bokwaongo invited trip; do not cold-gate'})

wb.save('leads/leads_50.xlsx')
print('xlsx patched')
