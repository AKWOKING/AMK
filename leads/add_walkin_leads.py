# -*- coding: utf-8 -*-
"""Register the Tue 15 Sep Buea walk-in batch (7 new leads + enrich 4 existing)."""
import openpyxl

wb = openpyxl.load_workbook('leads/leads_50.xlsx')
ws = wb['Leads 50']
hdr = [c.value for c in ws[1]]
C = {h: i + 1 for i, h in enumerate(hdr)}

# ---- enrich existing Buea rows ----
extra = {
    18: ' | Tue 15 walk-in Loop 2 (Great Soppo gate TBC); WA-check 652 075 229 + 675 533 321 Mon night; email pcss@yahoo.com; boarding hook + Presbyterian network angle; confessional principal-first.',
    19: ' | Tue 15 BACKUP walk-in if in quarter (Muea stretch swap).',
    20: ' | Tue 15 walk-in stop #2 Molyko; zero online contacts = gate visit solves it; capture principal WA + enrollment poster.',
    21: ' | Tue 15 BACKUP walk-in (technical streams angle).',
}
for r in range(2, ws.max_row + 1):
    lid = ws.cell(r, 1).value
    if lid in extra:
        c = ws.cell(r, C['Notes'])
        c.value = (c.value or '') + extra[lid]

new_leads = [
    {"ID": 27, "School": "St. Theresa International Bilingual Comprehensive College (STIBCCOL)",
     "City": "Buea (Molyko)", "Language": "EN/FR",
     "Facebook": "Page 100067292716772 (admissions-ongoing post Aug 2024)", "Website": "None found",
     "Website status": "D/E - no site", "Admissions activity": "Yes (FB admissions posts)",
     "Phone": "via FB page (not listed)", "WhatsApp": "verify on walk-in",
     "Decision maker": "Principal/proprietor (capture)",
     "Contact channel": "WALK-IN Tue 15 stop 1 + stibccol@yahoo.com",
     "Facilities": "Large GCE centre 11555 (~128 candidates 2024)", "Lead score": 14, "Priority": "A",
     "Notes": "Walk-In-Batch-2026-09-15. Pain: parents find only FB/yahoo, no website. Bilingual international comprehensive. Get WA + permission, preview by 4pm."},
    {"ID": 28, "School": "Summerset Bilingual College (SMBICOL)",
     "City": "Buea (Wokoko)", "Language": "EN/FR",
     "Website": "None found", "Website status": "D/E - no site",
     "Admissions activity": "Likely (verify on gate)",
     "Phone": "not listed", "WhatsApp": "verify on walk-in",
     "Decision maker": "Proprietor/principal (capture)",
     "Contact channel": "WALK-IN Tue 15 stop 4 + smbicol@yahoo.com",
     "Facilities": "~1,000 students / ~80 teachers (2021 paper); GCE centre 11447; founded 2000 lay private",
     "Lead score": 15, "Priority": "A",
     "Notes": "Walk-In-Batch-2026-09-15 FLAGSHIP stop: biggest enrollment of the day, zero web presence. Full 90s if allowed."},
    {"ID": 29, "School": "Saint Bernard High School (SBHS)",
     "City": "Buea (Molyko)", "Language": "EN",
     "Website": "None found", "Website status": "D/E - no site",
     "Admissions activity": "Yes - growth-stage",
     "Phone": "not listed", "WhatsApp": "verify on walk-in",
     "Decision maker": "Dr. Bernard Fomba (founder/proprietor, Oct 2020)",
     "Contact channel": "WALK-IN Tue 15 stop 3",
     "Facilities": "Forms 1-Upper 6th, general + commercial; ~247 at founding; GCE centre 12652",
     "Lead score": 12, "Priority": "B+",
     "Notes": "Walk-In-Batch-2026-09-15. Young school 2020 = hungry for growth, budget caution; founder decides alone."},
    {"ID": 30, "School": "NABESK Comprehensive College",
     "City": "Buea (Bonduma)", "Language": "EN",
     "Website": "None found", "Website status": "D/E - no site",
     "Phone": "not listed", "WhatsApp": "verify on walk-in",
     "Decision maker": "Principal (capture)",
     "Contact channel": "WALK-IN Tue 15 stop 5",
     "Facilities": "GCE centre 11853; 83.5% O-level 2020; hosts 400+ external candidates = strong reputation",
     "Lead score": 13, "Priority": "A-",
     "Notes": "Walk-In-Batch-2026-09-15. Hook: results buried in GCE PDFs, should be front page of own site."},
    {"ID": 31, "School": "Baird Memorial College",
     "City": "Buea (Bonduma)", "Language": "EN",
     "Website": "None found", "Website status": "D/E - no site",
     "Phone": "not listed", "WhatsApp": "verify on walk-in",
     "Decision maker": "Principal (capture)",
     "Contact channel": "WALK-IN Tue 15 stop 6 (light, same road as NABESK)",
     "Facilities": "GCE centre 11489; small (20-40 candidates)", "Lead score": 7, "Priority": "C+",
     "Notes": "Walk-In-Batch-2026-09-15. Card + 2 min; deepen only if interest shown."},
    {"ID": 32, "School": "Baptist Comprehensive College",
     "City": "Buea (Great Soppo)", "Language": "EN",
     "Facebook": "Page lists +237 679 65 07 07", "Website": "None found",
     "Website status": "D/E - no site",
     "Phone": "+237 679 65 07 07", "WhatsApp": "CHECK Mon night",
     "Decision maker": "Principal; Baptist education office above",
     "Contact channel": "WA-CHECK then WALK-IN Tue 15 stop 7; bgccbuea@yahoo.com",
     "Facilities": "GCE centre 11708", "Lead score": 12, "Priority": "B+",
     "Notes": "Walk-In-Batch-2026-09-15. Confessional: principal first, ask 'who else decides?', offer forwardable version. Contingent WA msg in walk-in doc."},
    {"ID": 33, "School": "St. Sylvester International College",
     "City": "Buea (Muea)", "Language": "EN/FR",
     "Website": "None found", "Website status": "D/E - no site",
     "Phone": "not listed", "WhatsApp": "verify on walk-in",
     "Decision maker": "Principal (capture)",
     "Contact channel": "WALK-IN Tue 15 STRETCH stop 10 (if ahead of schedule)",
     "Facilities": "GCE centre 12908; 48 candidates 2025; international private", "Lead score": 8, "Priority": "B-",
     "Notes": "Walk-In-Batch-2026-09-15. Swap for Frankfils/Marthlo depending on quarter. 'International' = fee capacity."},
]
for nl in new_leads:
    ws.append([nl.get(h) for h in hdr])

# Pipeline tab: next actions for the new seven
if 'Pipeline - 5 Stages' in wb.sheetnames:
    ps = wb['Pipeline - 5 Stages']
    h2 = [c.value for c in ps[1]]
    def pcol(name):
        for i, h in enumerate(h2, start=1):
            if h and name.lower() in str(h).lower():
                return i
        return None
    ci_name = pcol('lead') or 2
    ci_stage = pcol('stage') or 6
    ci_next = pcol('next') or 7
    for nl in new_leads:
        ps.append([None] * (ci_name - 1) + [nl['School']] + [None] * (ci_stage - ci_name - 1) +
                  ['1 -> 2 walk-in Tue 15', f"Tue 15 walk-in (sales/Walk-In-Batch-2026-09-15.md): meet principal, get WA, book preview review; concept within 24h on warm yes."])

wb.save('leads/leads_50.xlsx')
print('added', len(new_leads), 'leads; total', ws.max_row - 1)
