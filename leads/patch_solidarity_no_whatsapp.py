"""15 Sep: King confirmed clinic line 677 61 57 57 is NOT on WhatsApp.
Update CRM ID 36 channel/notes/follow-up only; no facts about people change."""
import openpyxl

PATH = "leads/leads_50.xlsx"
wb = openpyxl.load_workbook(PATH)
ws = wb["Leads 50"]
headers = [c.value for c in ws[1]]

def col(name):
    return headers.index(name) + 1

for row in ws.iter_rows(min_row=2):
    if row[col("ID") - 1].value == 36:
        row[col("WhatsApp") - 1].value = "677 61 57 57 NOT on WhatsApp (King 15 Sep); fallback admin 677 61 12 07 unverified; 691 63 29 41 = Dr Njang personal, HOLD"
        row[col("Contact channel") - 1].value = "WA blocked on clinic line -> check SHF-CIG admin 677 61 12 07; else Wed 16 Sep sealed envelope to 'current medical director or clinic manager'"
        row[col("Follow-up date") - 1].value = "2026-09-16"
        n = row[col("Notes") - 1].value or ""
        tag = " [15 Sep send morning: 677 61 57 57 not registered on WhatsApp per King; voice line only, never cold-call; MITOC now leads Tue 08:35; Solidarity -> admin-line check, else Wed sealed card]"
        if tag.strip() not in n:
            row[col("Notes") - 1].value = n + tag
wb.save(PATH)

# readback
wb = openpyxl.load_workbook(PATH)
ws = wb["Leads 50"]
H = [c.value for c in ws[1]]
for row in ws.iter_rows(min_row=2):
    d = {h: c.value for h, c in zip(H, row)}
    if d["ID"] == 36:
        for k in ("School", "Phone", "WhatsApp", "Contact channel", "Follow-up date"):
            print(k, "=>", d[k])
