# -*- coding: utf-8 -*-
"""Generates: AMK-Client-Kickoff-Kit.docx (EN+FR) and AMK-Internal-Delivery-Checklist.docx"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0F, 0x17, 0x2A)
AMBER = RGBColor(0xB4, 0x53, 0x09)
SLATE = RGBColor(0x47, 0x55, 0x69)
GREEN = RGBColor(0x05, 0x96, 0x69)

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = NAVY
    return doc

def title(doc, text, sub=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
    if sub:
        p2 = doc.add_paragraph()
        r2 = p2.add_run(sub)
        r2.font.size = Pt(10.5); r2.font.color.rgb = AMBER; r2.bold = True
    # amber rule
    pr = doc.add_paragraph()
    rr = pr.add_run("━" * 34)
    rr.font.color.rgb = AMBER; rr.font.size = Pt(8)

def h2(doc, text, color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(13.5); r.font.color.rgb = color

def p(doc, text, color=None, size=10.5, bold=False, italic=False, after=4):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(after)
    r = par.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r.font.color.rgb = color or SLATE
    return par

def item(doc, n, lead, detail, lead_fr=None):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run(f"{n}.  {lead}")
    r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run(detail)
    r2.font.size = Pt(10); r2.font.color.rgb = SLATE

def timeline(doc, rows):
    for i, (day, txt) in enumerate(rows):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
        r = par.add_run(f"  {day}   ")
        r.bold = True; r.font.size = Pt(10); r.font.color.rgb = AMBER
        r2 = par.add_run(txt)
        r2.font.size = Pt(10); r2.font.color.rgb = NAVY

# ============================================================
# DOCUMENT 1 — CLIENT KICKOFF KIT (EN + FR)
# ============================================================
doc = new_doc()
title(doc, "New School Website — Client Kickoff",
      "AMK – Web Development & Digital Solutions  ·  Douala · Yaoundé · Buea · Limbe")

p(doc, "Thank you for trusting AMK. This one-page checklist is everything we need to build and launch "
       "your school's website in 3–5 days. Send the items in any order — or better: book the 30-minute call "
       "and we'll fill it in together.", after=6)

h2(doc, "ENGLISH")

h2(doc, "How your 5 days will go", color=GREEN)
timeline(doc, [
    ("Day 0", "You sign and pay the 50% deposit. We send you this checklist. (You are here.)"),
    ("Day 0–1", "You send us items 1–8 below — photos and answers in this chat, or on a 30-minute call."),
    ("Day 1", "We confirm everything is in place and start building your bilingual site."),
    ("Day 1–2", "You receive the near-final version for one last check (your 48-hour review window)."),
    ("Day 3–5", "Domain, secure hosting, final checks — your school goes LIVE."),
    ("Day 6–36", "One month of free support + your 30-minute staff training."),
])

h2(doc, "The 8 things we need from you")
item(doc, 1, "Your crest or logo",
     "A PNG, PDF or JPG is fine — even a clear photo of the crest works, we clean it up. No crest yet? Tell us and we'll prepare one for your approval.")
item(doc, 2, "Your school colors",
     "Two or three colors. A photo of your uniform or building is enough, or simply tell us (e.g. 'green and gold').")
item(doc, 3, "10–15 photos",
     "Building/campus (2) · classrooms (3) · students learning (3) · sports or outdoor (2) · an event or award (2) · staff group (1, optional). "
     "Phone photos are perfectly fine — landscape (sideways) photos look best.")
item(doc, 4, "Your school story — 3 to 5 sentences",
     "When you were founded, by whom, what you are known for, and your motto if you have one. Write it the way you would tell a visitor — we polish it for the web.")
item(doc, 5, "Programs & key facts",
     "Stages offered (nursery / primary / secondary + ages) · language(s) of instruction · typical class size · accreditation · notable results or awards (last 2–3 years).")
item(doc, 6, "Contact details to publish",
     "Main phone number · WhatsApp number (WhatsApp Business if you have one — the site's form will use it) · email · full street address · office hours · "
     "Facebook / TikTok links if you have them.")
item(doc, 7, "Admissions & fees",
     "Your 3 admission steps (how a parent applies) · current fee schedule, or the words 'fees available on request — call us' · intake dates for the new academic year if known.")
item(doc, 8, "Domain & launch date",
     "Your preferred web address (e.g. yourschool.cm or yourschool.com — we check availability and set it up) · the date you'd like to go live "
     "(ideally before a term start or open day).")

h2(doc, "Two easy ways to send it")
p(doc, "A — Reply in this chat: send the photos and answers in any order, across as many messages as you need.", bold=False)
p(doc, "B — A 30-minute WhatsApp voice call (recommended): we ask, you answer, we fill everything in for you. No typing, no stress.")

h2(doc, "What you get", color=GREEN)
p(doc, "• One-page bilingual website — English & French, one tap to switch (standard on every AMK site)")
p(doc, "• WhatsApp button on every screen — parents reach you in one tap")
p(doc, "• Domain + secure (SSL) hosting, set up for you")
p(doc, "• 30-minute staff training so your own team can update news and photos")
p(doc, "• Live in 3–5 days · one month of free support after launch")
p(doc, "• Your domain and your files remain 100% yours")

h2(doc, "Your AMK contact")
p(doc, "Akwo King — Founder", bold=True, color=NAVY)
p(doc, "+237 677 789 631 (WhatsApp)  ·  hello@amk-cm.com")
p(doc, "We reply the same day, Monday to Saturday.", italic=True)

# ---- divider ----
pd = doc.add_paragraph()
rd = pd.add_run("━" * 34)
rd.font.color.rgb = AMBER; rd.font.size = Pt(8)
p(doc, "", after=0)

h2(doc, "FRANÇAIS")

p(doc, "Merci de faire confiance à AMK. Cette check-list d'une page est tout ce qu'il nous faut pour construire "
       "et lancer le site de votre école en 3 à 5 jours. Envoyez les éléments dans l'ordre que vous voulez — "
       "ou mieux : réservez l'appel de 30 minutes et nous la remplissons ensemble.", after=6)

h2(doc, "Comment se passent vos 5 jours", color=GREEN)
timeline(doc, [
    ("Jour 0", "Vous signez et payez l'acompte de 50 %. Nous vous envoyons cette check-list. (Vous êtes ici.)"),
    ("Jour 0–1", "Vous nous envoyez les éléments 1 à 8 ci-dessous — photos et réponses dans cette conversation, ou lors d'un appel de 30 minutes."),
    ("Jour 1", "Nous confirmons que tout est en place et nous démarrons la construction de votre site bilingue."),
    ("Jour 1–2", "Vous recevez la version quasi finale pour un dernier contrôle (votre fenêtre de 48 heures)."),
    ("Jour 3–5", "Domaine, hébergement sécurisé, vérifications finales — votre école est EN LIGNE."),
    ("Jour 6–36", "Un mois de support gratuit + votre formation du personnel de 30 minutes."),
])

h2(doc, "Les 8 éléments dont nous avons besoin")
item(doc, 1, "Vos armoiries ou votre logo",
     "PNG, PDF ou JPG — même une photo nette des armoiries suffit, nous la nettoyons. Pas encore d'armoiries ? Dites-le nous, nous en préparons une pour votre validation.")
item(doc, 2, "Les couleurs de votre école",
     "Deux ou trois couleurs. Une photo de votre uniforme ou de votre bâtiment suffit, ou dites-le simplement (ex. « vert et or »).")
item(doc, 3, "10 à 15 photos",
     "Bâtiment/campus (2) · salles de classe (3) · élèves en apprentissage (3) · sport ou extérieur (2) · un événement ou distinction (2) · groupe du personnel (1, facultatif). "
     "Les photos de téléphone sont parfaitement acceptées — les photos en paysage (à l'horizontale) sont les plus belles.")
item(doc, 4, "L'histoire de votre école — 3 à 5 phrases",
     "Quand vous avez été fondé(e), par qui, ce pour quoi vous êtes connu(e), et votre devise si vous en avez une. Écrivez comme vous le diriez à un visiteur — nous soignons le texte pour le web.")
item(doc, 5, "Programmes & faits clés",
     "Niveaux proposés (maternelle / primaire / secondaire + âges) · langue(s) d'enseignement · taille habituelle des classes · accréditation · résultats ou distinctions notables (2–3 dernières années).")
item(doc, 6, "Coordonnées à publier",
     "Numéro principal · numéro WhatsApp (WhatsApp Business si vous en avez un — le formulaire du site l'utilisera) · e-mail · adresse complète · heures d'ouverture · "
     "liens Facebook / TikTok si vous en avez.")
item(doc, 7, "Admissions & frais",
     "Vos 3 étapes d'admission (comment un parent postule) · le barème actuel des frais, ou la mention « frais sur demande — appelez-nous » · dates de rentrée pour la nouvelle année scolaire si connues.")
item(doc, 8, "Domaine & date de lancement",
     "Votre adresse web préférée (ex. votre-ecole.cm ou votre-ecole.com — nous vérifions la disponibilité et nous l'installons) · la date à laquelle vous souhaitez être en ligne "
     "(idéalement avant une rentrée ou des portes ouvertes).")

h2(doc, "Deux façons faciles de nous les envoyer")
p(doc, "A — Répondez dans cette conversation : envoyez photos et réponses dans l'ordre que vous voulez, en autant de messages que nécessaire.")
p(doc, "B — Un appel WhatsApp de 30 minutes (recommandé) : nous posons les questions, vous répondez, nous remplissons tout pour vous. Pas de frappe, pas de stress.")

h2(doc, "Ce que vous recevez", color=GREEN)
p(doc, "• Un site web bilingue d'une page — anglais et français, un geste pour basculer (standard sur tous les sites AMK)")
p(doc, "• Un bouton WhatsApp sur chaque écran — les parents vous joignent en un clic")
p(doc, "• Domaine + hébergement sécurisé (SSL), installés pour vous")
p(doc, "• Une formation de 30 minutes pour que votre équipe mette à jour les actualités et photos")
p(doc, "• En ligne en 3–5 jours · un mois de support gratuit après le lancement")
p(doc, "• Votre domaine et vos fichiers vous appartiennent à 100 %")

h2(doc, "Votre contact AMK")
p(doc, "Akwo King — Fondateur", bold=True, color=NAVY)
p(doc, "+237 677 789 631 (WhatsApp)  ·  hello@amk-cm.com")
p(doc, "Nous répondons le jour même, du lundi au samedi.", italic=True)

doc.save("/home/user/agency/sales/AMK-Client-Kickoff-Kit.docx")
print("saved kickoff kit")

# ============================================================
# DOCUMENT 2 — INTERNAL DELIVERY CHECKLIST
# ============================================================
d2 = new_doc()
title(d2, "Internal Delivery Checklist",
      "One per signed client · keep in CRM notes + this folder · DO NOT forward to the client")

h2(d2, "Stage 4 → 5 handoff (Day 0)")
for t in [
    "□ 50% deposit received (MoMo/bank) — receipt sent to client, logged in CRM Pipeline sheet",
    "□ Kickoff Kit (EN or FR per school) forwarded with the 2-line cover message",
    "□ Client contact confirmed: who sends content, who gives final approval",
]:
    p(d2, t, color=NAVY, after=3)

h2(d2, "Content collection (Day 0–1)")
for t in [
    "□ Items 1–8 received (or 30-min call done) — missing items chased once, same day",
    "□ Crest cleaned/vectorized if needed",
    "□ Photos sorted: landscape priority, 10+ usable",
    "□ Fact sheet written up (story, programs, results, contacts, fees) — the source of truth for the build",
]:
    p(d2, t, color=NAVY, after=3)

h2(d2, "Build (Day 1–2)")
for t in [
    "□ Correct concept template chosen (academy / nursery-primary / secondary) and personalized: name, crest, colors, facts, CTAs on THEIR number",
    "□ EN|FR toggle tested in both directions — no untranslated leftovers",
    "□ Mobile QA: iPhone + Android, on 3G (target: under 3s first paint)",
    "□ All WhatsApp links tested (wa.me opens with the right pre-filled text)",
    "□ Form tested: payload logs to console + routes to the right number",
    "□ SEO: meta title <60 chars, description <155, JSON-LD filled with real school data (NO invented ratings)",
    "□ AMK credit line present (small, footer) — the agency badge, not the agency brand",
]:
    p(d2, t, color=NAVY, after=3)

h2(d2, "Launch (Day 3–5)")
for t in [
    "□ Domain registered in the CLIENT'S name — AMK manages DNS/hosting (trust + ownership)",
    "□ Hosting: files deployed (Vercel/Netlify), SSL active, DNS pointed",
    "□ Final check on the live URL (both languages) before telling the client",
    "□ Launch message to client with the live link + 50% balance request",
    "□ 30-min staff training scheduled and done (WhatsApp call; record the 3 most-asked updates)",
]:
    p(d2, t, color=NAVY, after=3)

h2(d2, "Proof & flywheel (launch week → Day 30)")
for t in [
    "□ Testimonial requested (one honest line from the director — screenshot of it is fine)",
    "□ 2–3 photos + the quote → case study → AMK site Proof section (replaces a 'slot' card)",
    "□ Referral asked explicitly: 'Which other school should we build for next?'",
    "□ CRM: stage 5 closed, next touch date = Day 30 support check-in",
    "□ Upsell menu offered once, softly: results page · e-learning/registration · other campuses · maintenance plan",
]:
    p(d2, t, color=NAVY, after=3)

h2(d2, "Rules of the road", color=GREEN)
for t in [
    "• The 24h preview promise runs from the FIRST message — the clock never waits for us.",
    "• Client content is the #1 delay risk — the 30-min call beats a missing photo every time.",
    "• Never invent facts on a client site (stats, ratings, dates). Placeholder → confirm → publish.",
    "• One WhatsApp message = one deliverable update. Daily until launch, even if the update is 'on track'.",
]:
    p(d2, t, color=NAVY, after=3)

d2.save("/home/user/agency/sales/AMK-Internal-Delivery-Checklist.docx")
print("saved delivery checklist")
