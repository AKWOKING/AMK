#!/usr/bin/env python3
"""AMK — petit constructeur xlsx propre à l'agence (openpyxl, MIT).

Né de l'audit du 20/10 (anthropics/skills) : les skills documentaires amont sont
« source-available » avec interdiction d'extraction/copie hors des Services Anthropic —
on ne les installe PAS. On garde la MÉTHODE (fonte professionnelle, zéro erreur de
formule, sources citées dans le fichier) et on l'implémente avec une pile libre.

Usage : python3 tools/docs/xlsx_build.py <entree.json> <sortie.xlsx> --title "..." --source "..."
Format entree.json : {"columns": [["Code", 10], ...], "rows": [[...], ...], "meta": "..."}
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Arial"
NAVY = "0A1220"
CREAM = "F6EBE0"
GREY = "EEF1F5"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("entree")
    ap.add_argument("sortie")
    ap.add_argument("--title", default="")
    ap.add_argument("--source", default="")
    args = ap.parse_args()

    data = json.loads(Path(args.entree).read_text(encoding="utf-8"))
    cols, rows = data["columns"], data["rows"]
    widths = [c[1] if isinstance(c, list) else 18 for c in cols]
    headers = [c[0] if isinstance(c, list) else c for c in cols]

    wb = Workbook()
    ws = wb.active
    ws.title = data.get("sheet", "Liste")
    line = 1
    if args.title:
        ws.cell(row=1, column=1, value=args.title)
        ws.cell(row=1, column=1).font = Font(name=FONT, bold=True, size=14, color=NAVY)
        line = 2
    if data.get("meta") or args.source:
        line += 1
        ws.cell(row=line, column=1, value=" · ".join(x for x in (data.get("meta"), args.source) if x))
        ws.cell(row=line, column=1).font = Font(name=FONT, italic=True, size=9, color="6E5F4E")
        line += 1
    line += 1

    head_row = line
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=head_row, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, size=10, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=CREAM)
        cell.border = Border(bottom=Side(style="thin", color=NAVY))
        cell.alignment = Alignment(vertical="center")
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w

    for r, row in enumerate(rows, head_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            if (r - head_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=GREY)
        ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)

    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    Path(args.sortie).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.sortie)
    print(f"ok : {args.sortie} — {len(rows)} lignes × {len(headers)} colonnes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
